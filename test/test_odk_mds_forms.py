"""The Crop Sown ODK forms start with the two identity questions and take their lists from Master Data
(odk/tools/build_mds_forms.py)."""
import csv
import importlib.util
import json
import re
from pathlib import Path

import openpyxl
import pytest

REPO = Path(__file__).resolve().parents[1]
spec = importlib.util.spec_from_file_location("build_mds_forms", REPO / "odk/tools/build_mds_forms.py")
builder = importlib.util.module_from_spec(spec)
spec.loader.exec_module(builder)


def mds_list(attribute_id, codes, **attributes):
    return {"attribute_id": attribute_id, "values": [
        {"value_id": c, "value_code": c, "value_display": c.replace("_", " ").title(), "sort_order": i + 1,
         "attributes": attributes.get(c, {})} for i, c in enumerate(codes)]}


@pytest.fixture
def pack(tmp_path):
    root = tmp_path / "ETH-TEST"
    (root / "domains/agriculture").mkdir(parents=True)
    lists = {
        "CROP_SEASON": mds_list("CROP_SEASON", ["SEASON_MEHER", "SEASON_BELG", "SEASON_IRRIGATION"]),
        "AGRO_ECOLOGICAL_ZONE": mds_list("AGRO_ECOLOGICAL_ZONE", ["AEZ_KOLLA", "AEZ_DEGA"]),
        "CROP_CATEGORY": mds_list("CROP_CATEGORY", ["CROPCAT_CEREAL", "CROPCAT_OILSEED"]),
        "CROP_COMMODITY": mds_list("CROP_COMMODITY", ["CROP_TEFF", "CROP_SESAME", "CROP_ALFALFA"],
                                   CROP_TEFF={"category": "CROPCAT_CEREAL"}, CROP_SESAME={"category": "CROPCAT_OILSEED"}),
        "SEED_VARIETY": mds_list("SEED_VARIETY", ["VAR_TEFF_QUNCHO", "VAR_TEFF_BOSET"]),
    }
    for attribute_id, data in lists.items():
        (root / "domains/agriculture" / f"{attribute_id.lower()}.json").write_text(json.dumps(data), encoding="utf-8")
    (root / "levels.json").write_text(json.dumps([
        {"level_id": "l0", "level_mnemonic": "country"}, {"level_id": "l1", "level_mnemonic": "region"},
        {"level_id": "l2", "level_mnemonic": "zone"}, {"level_id": "l3", "level_mnemonic": "woreda"},
        {"level_id": "l4", "level_mnemonic": "kebele"}]), encoding="utf-8")
    units = [{"level_value_id": "ET", "level_id": "l0", "display_name": "Ethiopia", "parent_level_value_id": None}]
    for r, name in (("01", "Tigray"), ("03", "Amhara")):
        units += [{"level_value_id": f"ET{r}", "level_id": "l1", "display_name": name, "parent_level_value_id": "ET"},
                  {"level_value_id": f"ET{r}01", "level_id": "l2", "display_name": f"Zone {r}", "parent_level_value_id": f"ET{r}"},
                  {"level_value_id": f"ET{r}0101", "level_id": "l3", "display_name": f"Woreda {r}", "parent_level_value_id": f"ET{r}01"},
                  {"level_value_id": f"ET{r}010101", "level_id": "l4", "display_name": f"Kebele {r}", "parent_level_value_id": f"ET{r}0101"}]
    (root / "values.json").write_text(json.dumps(units), encoding="utf-8")
    return root


def sheet(book, title):
    rows = list(book[title].iter_rows(values_only=True))
    return [dict(zip(rows[0], r)) for r in rows[1:] if any(r)]


def questions(rows):
    """The rows a person answers: not structure, not calculates, not notes."""
    skip = ("begin_group", "end_group", "begin_repeat", "end_repeat", "calculate", "note")
    return [r for r in rows if (r["type"] or "").split()[0] not in skip]


def first_group(survey):
    assert survey[0]["type"] == "begin_group" and survey[0]["name"] == "farmer_identity"
    end = next(i for i, r in enumerate(survey) if r["type"] == "end_group")
    return survey[1:end]


def check_identity(survey, form_id):
    inside = first_group(survey)
    asked = questions(inside)
    assert [r["name"] for r in asked] == ["farmer_id", "fayda_id_entered"], form_id
    assert all(r["required"] == "yes" for r in asked), form_id
    farmer, fayda = asked
    assert farmer["constraint"] == "regex(., '^FR-[0-9]{12}$')"
    assert fayda["constraint"] == "regex(translate(., ' ', ''), '^([0-9]{12}|[0-9]{16})$')"
    calculates = [r for r in inside if r["type"] == "calculate"]
    # the transform reads fayda_id: the typed value without its spaces
    assert [(r["name"], r["calculation"]) for r in calculates] == [("fayda_id", "translate(${fayda_id_entered}, ' ', '')")]
    for r in (farmer, fayda):
        for text in (r["label"], r["hint"], r["required_message"], r["constraint_message"]):
            assert text and "_" not in text, (form_id, text)


def check_lists(survey, choices, files_dir, form_id):
    used = builder.used_lists(survey)
    accounted = set(builder.INLINE_LISTS) | builder.YES_NO | set(builder.NOT_MASTER_DATA)
    assert used <= accounted, (form_id, sorted(used - accounted))
    assert {r["list_name"] for r in choices} == used, form_id  # no stray or missing list
    assert not used & set(builder.REPLACES), form_id  # the hard-coded season and zone lists are gone
    for r in survey:
        parts = (r["type"] or "").split()
        if parts[0] == "select_one_from_file":
            assert parts[1] in builder.FILE_LISTS or parts[1] in {f"{level}.csv" for level in builder.LEVELS}, (form_id, parts[1])
            assert (files_dir / parts[1]).exists(), (form_id, parts[1])


def test_each_form_has_the_two_identity_questions_and_master_data_lists(pack, tmp_path):
    root = pack
    out, media = tmp_path / "forms", tmp_path / "media"
    builder.build(root, None, "test-version", builder.SOURCE, out, media)
    lists = builder.read_lists(root / "domains")
    for form_id in builder.FORM_IDS:
        book = openpyxl.load_workbook(out / f"{form_id}.xlsx")
        survey, choices = sheet(book, "survey"), sheet(book, "choices")
        assert sheet(book, "settings")[0]["version"] == "test-version"
        assert sheet(book, "settings")[0]["form_id"] == form_id
        check_identity(survey, form_id)
        check_lists(survey, choices, media, form_id)
        for name, attribute in builder.INLINE_LISTS.items():
            rows = [r for r in choices if r["list_name"] == name]
            if rows:
                assert [r["name"] for r in rows] == [v["value_code"] for v in builder.ordered(lists[attribute]["values"])], (form_id, name)
                assert all(r["label"] for r in rows)
        # crop year and season sit in the group the transform reads, not in the identity group
        assert any(r["name"] == "production_season_info" for r in survey), form_id
        assert {"crop_year", "production_season"} <= {r["name"] for r in survey}, form_id


def test_media_files_are_master_data_codes(pack, tmp_path):
    root = pack
    media = tmp_path / "media"
    builder.build(root, None, "test-version", builder.SOURCE, tmp_path / "forms", media)

    def read(name):
        with (media / name).open(encoding="utf-8", newline="") as handle:
            return list(csv.DictReader(handle))

    crops = read("crop_name.csv")
    assert [r["name"] for r in crops] == ["CROP_TEFF", "CROP_SESAME", "CROP_ALFALFA"]
    assert [r["category"] for r in crops] == ["CROPCAT_CEREAL", "CROPCAT_OILSEED", ""]
    assert crops[0]["category_label"] == "Cropcat Cereal"
    assert not (media / "crop_variety.csv").exists()  # one variety list: SEED_VARIETY
    assert [r["name"] for r in read("seed_variety.csv")] == ["VAR_TEFF_QUNCHO", "VAR_TEFF_BOSET"]
    assert [r["name"] for r in read("region.csv")] == ["ET01", "ET03"]
    assert [(r["name"], r["parent"]) for r in read("zone.csv")] == [("ET0101", "ET01"), ("ET0301", "ET03")]
    assert [(r["name"], r["parent"]) for r in read("kebele.csv")] == [("ET01010101", "ET010101"), ("ET03010101", "ET030101")]


def test_a_list_nobody_accounted_for_stops_the_build(pack, tmp_path, monkeypatch):
    root = pack
    monkeypatch.delitem(builder.NOT_MASTER_DATA, "plot_cat")
    with pytest.raises(SystemExit, match="plot_cat"):
        builder.build(root, None, "test-version", builder.SOURCE, tmp_path / "forms", tmp_path / "media")


@pytest.mark.parametrize("form_id", builder.FORM_IDS)
def test_committed_workbooks(form_id):
    """The workbooks and CSVs in the repository: identity, lists, media."""
    book = openpyxl.load_workbook(REPO / f"odk/forms/{form_id}.xlsx")
    survey, choices = sheet(book, "survey"), sheet(book, "choices")
    settings = sheet(book, "settings")[0]
    assert settings["form_id"] == form_id and re.fullmatch(r"\d{4}\.\d{2}\.\d{2}(\.\d+)?", str(settings["version"]))
    check_identity(survey, form_id)
    check_lists(survey, choices, REPO / "odk/media/mds", form_id)
    names = {}
    for r in choices:
        names.setdefault(r["list_name"], []).append(r["name"])
    if "crop_season" in names:
        assert set(names["crop_season"]) == {"SEASON_MEHER", "SEASON_BELG", "SEASON_IRRIGATION"}
    for r in survey:
        kind = (r["type"] or "").split()
        if kind[0] == "select_one_from_file" and kind[1] in ("zone.csv", "woreda.csv", "kebele.csv"):
            assert (r["choice_filter"] or "").startswith("parent="), (form_id, r["name"])


@pytest.mark.parametrize("form_id", ("crop_sown_registry_plan", "crop_sown_registry_prep"))
def test_the_forms_ask_for_one_variety_called_variety(pack, tmp_path, form_id):
    builder.build(pack, None, "test-version", builder.SOURCE, tmp_path / "forms", tmp_path / "media")
    survey = sheet(openpyxl.load_workbook(tmp_path / "forms" / f"{form_id}.xlsx"), "survey")
    varieties = [r for r in survey if "variety" in (r["name"] or "")]
    assert [(r["name"], r["label"], r["type"]) for r in varieties] == [("seed_variety", "Variety", "select_one_from_file seed_variety.csv")]
